"""
Формирование итогового Excel-отчёта сверки БУ-ПУ.

Лист «Свод»: сводная таблица с заголовком, 14 строк показателей,
  итоговыми суммами по БУ и ПУ и расхождением (сумма + конкретные даты).
Лист «Детализация»: строки с расхождениями по каждой дате.

Структура заголовка (строки 1-3):
  Строка 1: «Сверка БУ-ПУ. <Период>» (A1:E1 merged)
  Строка 2: «Виды движений» (A2:A3) | «Бухгалтерский учёт» (B2:B3) |
            «Персонифицированный учёт» (C2:C3) | «Расхождение» (D2:E2)
  Строка 3: (под merge) | (под merge) | (под merge) | Сумма | Дата
  Строки 4-17: данные (14 показателей)
"""
from __future__ import annotations

import logging
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from npf_recon.models import ReconRow
from npf_recon.normalize import format_amount, format_date

logger = logging.getLogger(__name__)

# Цвета
_HEADER_FILL = PatternFill("solid", fgColor="D9E1F2")
_DISCREP_FILL = PatternFill("solid", fgColor="FFE0E0")

_THIN = Side(style="thin")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _apply_border(cell) -> None:
    cell.border = _BORDER


def _format_dates(dates: list[date]) -> str:
    """
    Форматирует список дат расхождений для ячейки отчёта.

    Перечисляет ВСЕ конкретные даты, в которые есть расхождение, через «, »
    (диапазоны не используются — нужны именно конкретные даты).
    """
    if not dates:
        return ""
    return ", ".join(format_date(d) for d in sorted(dates))


def write_report(
    rows: list[ReconRow],
    output_path: Path,
    period_label: str,
    eps: float = 0.005,
) -> None:
    """
    Записывает итоговый отчёт в Excel-файл.

    :param rows:         список из 14 ReconRow от reconcile().
    :param output_path:  путь к выходному файлу.
    :param period_label: метка периода (напр. «Февраль 2026»).
    :param eps:          допуск: расхождения |d| <= eps не выделяются.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    wb = openpyxl.Workbook()
    _write_svod(wb, rows, period_label, eps)
    _write_details(wb, rows)

    wb.save(output_path)
    logger.info("Отчёт сохранён: %s", output_path)


# ------------------------------------------------------------------
# Лист «Свод»
# ------------------------------------------------------------------

def _write_svod(
    wb: openpyxl.Workbook,
    rows: list[ReconRow],
    period_label: str,
    eps: float,
) -> None:
    ws = wb.active
    ws.title = "Свод"
    bold = Font(bold=True)

    def _head(ref: str, text: str) -> None:
        cell = ws[ref]
        cell.value = text
        cell.font = bold
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.fill = _HEADER_FILL
        _apply_border(cell)

    # --- Строка 1: заголовок периода (A1:E1) ---
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_cell.value = f"Сверка БУ-ПУ. {period_label}"
    title_cell.font = Font(bold=True, size=13)
    title_cell.alignment = Alignment(horizontal="center", vertical="center")

    # --- Строки 2-3: шапка таблицы ---
    # БУ и ПУ — по одной колонке итоговой суммы; Расхождение — Сумма + Дата.
    ws.merge_cells("A2:A3")
    _head("A2", "Виды движений")
    _apply_border(ws["A3"])

    ws.merge_cells("B2:B3")
    _head("B2", "Бухгалтерский учёт")
    _apply_border(ws["B3"])

    ws.merge_cells("C2:C3")
    _head("C2", "Персонифицированный учёт")
    _apply_border(ws["C3"])

    ws.merge_cells("D2:E2")
    _head("D2", "Расхождение")
    _apply_border(ws["E2"])
    _head("D3", "Сумма")
    _head("E3", "Дата")

    # --- Строки данных (строки 4-17) ---
    for data_row_idx, recon in enumerate(rows, start=4):
        bu = recon.bu
        pu = recon.pu
        # Расхождение по итогу периода (сумма БУ − ПУ за месяц)
        has_total_discrep = (
            recon.diff_total is not None and abs(recon.diff_total) > eps
        )
        # Расхождение хотя бы по одному дню — даже если месячные итоги совпали
        # (взаимно компенсирующиеся отклонения по дням).
        has_any_discrep = has_total_discrep or bool(recon.diff_dates)

        row_data = [
            # (значение, колонка 1-based, выравнивание)
            (recon.indicator, 1, "left"),
            (format_amount(bu.total) if bu else "", 2, "right"),
            (format_amount(pu.total) if pu else "", 3, "right"),
            (format_amount(recon.diff_total) if has_total_discrep else "", 4, "right"),
            (_format_dates(recon.diff_dates) if recon.diff_dates else "", 5, "left"),
        ]

        for value, col_idx, align in row_data:
            cell = ws.cell(row=data_row_idx, column=col_idx, value=value)
            cell.alignment = Alignment(
                horizontal=align, vertical="top", wrap_text=True
            )
            _apply_border(cell)
            if has_any_discrep and col_idx in (4, 5):
                cell.fill = _DISCREP_FILL

    # --- Ширины колонок ---
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 16
    ws.column_dimensions["E"].width = 40

    ws.freeze_panes = "A4"


# ------------------------------------------------------------------
# Лист «Детализация»
# ------------------------------------------------------------------

def _write_details(
    wb: openpyxl.Workbook,
    rows: list[ReconRow],
) -> None:
    ws = wb.create_sheet("Детализация")
    bold = Font(bold=True)

    headers = ["Показатель", "Дата", "Сумма БУ", "Сумма ПУ", "Разница"]
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = bold
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
        _apply_border(cell)

    detail_row = 2
    for recon in rows:
        if not recon.diff_dates:
            continue
        for d in sorted(recon.diff_dates):
            bu_val = recon.bu.daily.get(d, 0.0) if recon.bu else 0.0
            pu_val = recon.pu.daily.get(d, 0.0) if recon.pu else 0.0
            diff_val = recon.diff_by_date.get(d)

            row_data = [
                recon.indicator,
                format_date(d),
                format_amount(bu_val),
                format_amount(pu_val),
                format_amount(diff_val),
            ]
            for col_idx, value in enumerate(row_data, start=1):
                cell = ws.cell(row=detail_row, column=col_idx, value=value)
                cell.alignment = Alignment(
                    horizontal="left" if col_idx <= 2 else "right"
                )
                _apply_border(cell)
                if col_idx == 5:
                    cell.fill = _DISCREP_FILL
            detail_row += 1

    ws.column_dimensions["A"].width = 42
    ws.column_dimensions["B"].width = 14
    for col_letter in ("C", "D", "E"):
        ws.column_dimensions[col_letter].width = 16

    ws.freeze_panes = "A2"
