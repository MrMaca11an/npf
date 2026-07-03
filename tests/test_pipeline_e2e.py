"""
Сквозной (e2e) тест pipeline: генерация данных → run → проверка отчёта.
"""
from __future__ import annotations

import os
import sys
from pathlib import Path

import openpyxl
import pytest


def _generate_in(data_dir: Path) -> None:
    """Генерирует тестовые данные в указанную директорию."""
    # Импортируем generate_sample_data, передавая data_dir
    scripts_dir = Path(__file__).resolve().parent.parent / "scripts"
    sys.path.insert(0, str(scripts_dir))
    import generate_sample_data
    generate_sample_data.generate(data_dir=data_dir)


@pytest.fixture
def temp_data_dir(tmp_path: Path):
    """
    Фикстура: создаёт временную директорию с тестовыми данными и
    настраивает NPF_DATA_DIR, чтобы pipeline использовал её.
    """
    _generate_in(tmp_path)

    # Переопределяем окружение
    old_env = os.environ.get("NPF_DATA_DIR")
    os.environ["NPF_DATA_DIR"] = str(tmp_path)
    yield tmp_path
    # Восстанавливаем
    if old_env is None:
        os.environ.pop("NPF_DATA_DIR", None)
    else:
        os.environ["NPF_DATA_DIR"] = old_env


@pytest.fixture
def output_file(tmp_path: Path) -> Path:
    """Путь к выходному файлу в temp_path."""
    return tmp_path / "output" / "Сверка_БУ_ПУ.xlsx"


def _run_pipeline(data_dir: Path, output_path: Path) -> int:
    """Запускает pipeline с временными путями."""
    from npf_recon.config import Config
    from npf_recon.pipeline import run

    config = Config(
        pu_dir=data_dir / "ПУ",
        bu_dir=data_dir / "БУ",
        output_file=output_path,
        period_label="Февраль 2026",
        eps=0.005,
    )
    return run(config)


# ------------------------------------------------------------------
# e2e тест
# ------------------------------------------------------------------

class TestPipelineE2E:
    """Сквозной тест от генерации данных до проверки отчёта."""

    def test_pipeline_runs_successfully(self, temp_data_dir: Path, output_file: Path):
        """Pipeline завершается с кодом 0."""
        exit_code = _run_pipeline(temp_data_dir, output_file)
        assert exit_code == 0

    def test_output_file_created(self, temp_data_dir: Path, output_file: Path):
        """Файл отчёта создаётся."""
        _run_pipeline(temp_data_dir, output_file)
        assert output_file.exists()

    def test_svod_sheet_exists(self, temp_data_dir: Path, output_file: Path):
        """Лист «Свод» присутствует."""
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        assert "Свод" in wb.sheetnames

    def test_svod_has_14_data_rows(self, temp_data_dir: Path, output_file: Path):
        """На листе «Свод» ровно 14 строк данных (плюс 3 строки шапки)."""
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        ws = wb["Свод"]
        # Строки 4..17 — данные (3 строки шапки + 14 показателей)
        data_rows = []
        for row in ws.iter_rows(min_row=4, max_row=17, min_col=1, max_col=1):
            cell = row[0]
            val = cell.value
            if val:
                data_rows.append(str(val))
        assert len(data_rows) == 14

    def test_svod_indicators_in_order(self, temp_data_dir: Path, output_file: Path):
        """Показатели в «Свод» расположены в фиксированном порядке."""
        from config.mappings import REPORT_INDICATORS
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        ws = wb["Свод"]
        actual = []
        for row in ws.iter_rows(min_row=4, max_row=17, min_col=1, max_col=1):
            cell = row[0]
            if cell.value:
                actual.append(str(cell.value))
        assert actual == REPORT_INDICATORS

    def test_strahovoy_rezerv_has_discrepancy(self, temp_data_dir: Path, output_file: Path):
        """«Страховой резерв» показывает расхождение на дату 17.02.2026."""
        from config.mappings import REPORT_INDICATORS
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        ws = wb["Свод"]

        # Находим строку «Страховой резерв»
        sr_row_idx = None
        for row in ws.iter_rows(min_row=4, max_row=17):
            if row[0].value == "Страховой резерв":
                sr_row_idx = row[0].row
                break

        assert sr_row_idx is not None, "Строка 'Страховой резерв' не найдена в Свод"

        # Колонка D (4) — «Расхождение Сумма»
        discrep_sum = ws.cell(row=sr_row_idx, column=4).value
        assert discrep_sum, "Расхождение по Страховому резерву должно быть непустым"

        # Колонка E (5) — «Расхождение Дата»
        discrep_date = ws.cell(row=sr_row_idx, column=5).value
        assert discrep_date is not None and "17.02.2026" in str(discrep_date), \
            f"Ожидалась дата 17.02.2026 в колонке расхождений, получено: {discrep_date}"

    def test_id_row_is_empty(self, temp_data_dir: Path, output_file: Path):
        """Строка «ИД» не имеет данных (нет файлов)."""
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        ws = wb["Свод"]

        for row in ws.iter_rows(min_row=4, max_row=17):
            if row[0].value == "ИД":
                bu_sum = row[1].value   # B — БУ Сумма
                pu_sum = row[2].value   # C — ПУ Сумма
                assert not bu_sum, f"БУ Сумма для ИД должна быть пустой, получено: {bu_sum}"
                assert not pu_sum, f"ПУ Сумма для ИД должна быть пустой, получено: {pu_sum}"
                return
        pytest.fail("Строка 'ИД' не найдена в Свод")

    def test_details_sheet_exists(self, temp_data_dir: Path, output_file: Path):
        """Лист «Детализация» присутствует."""
        _run_pipeline(temp_data_dir, output_file)
        wb = openpyxl.load_workbook(output_file)
        assert "Детализация" in wb.sheetnames

    @staticmethod
    def _details_map(ws) -> dict:
        """
        Восстанавливает группированную «Детализацию» в {показатель: {дата: разница}}.

        Учитывает новую структуру: строка-заголовок с названием показателя сверху,
        под ней — дневные строки без названия (колонка A пуста).
        """
        result: dict[str, dict[str, str]] = {}
        current = None
        for row in ws.iter_rows(min_row=2, values_only=True):
            name, date_str, bu_sum, pu_sum, diff = row
            if name and not any((date_str, bu_sum, pu_sum, diff)):
                current = name            # строка-заголовок группы
                result.setdefault(current, {})
            elif current and date_str:
                result[current][date_str] = diff
        return result

    def test_details_has_discrepancy_row(self, temp_data_dir: Path, output_file: Path):
        """«Детализация» содержит расхождение по Страховому резерву на 17.02.2026."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Детализация"]
        details = self._details_map(ws)
        assert "17.02.2026" in details.get("Страховой резерв", {})
        assert "15 000,00" in str(details["Страховой резерв"]["17.02.2026"])

    def test_pv_fl_discrepancy_in_details(self, temp_data_dir: Path, output_file: Path):
        """«Детализация» содержит расхождение по Пенсионным взносам ФЛ на 28.02.2026."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Детализация"]
        details = self._details_map(ws)
        assert "28.02.2026" in details.get("Пенсионные взносы ФЛ", {})

    @staticmethod
    def _svod_row(ws, indicator: str):
        """Возвращает кортеж значений строки «Свод» по названию показателя."""
        for row in ws.iter_rows(min_row=4, max_row=17, values_only=True):
            if row[0] == indicator:
                return row
        return None

    def test_cel_ul_daily_only_discrepancy(self, temp_data_dir: Path, output_file: Path):
        """Целевые взносы ЮЛ: итоги равны, но дневные расхождения отражены в колонке дат."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Свод"]
        row = self._svod_row(ws, "Целевые взносы ЮЛ")
        assert row is not None
        bu_sum, pu_sum = row[1], row[2]       # B=БУ, C=ПУ
        discrep_sum, discrep_dates = row[3], row[4]  # D=Расх.Сумма, E=Расх.Дата
        # Итоги БУ и ПУ совпадают → «Сумма расхождения» пустая (0 — не расхождение),
        # но конкретные даты дневных отличий перечислены.
        assert bu_sum == pu_sum
        assert not discrep_sum
        assert discrep_dates and "10.02.2026" in str(discrep_dates) and "24.02.2026" in str(discrep_dates)

    def test_svod_zero_diff_cell_not_highlighted(self, temp_data_dir: Path, output_file: Path):
        """«Сумма» пустая при нулевой разнице не подсвечивается красным."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Свод"]
        for row in ws.iter_rows(min_row=4, max_row=17):
            if row[0].value == "Целевые взносы ЮЛ":
                sum_cell = row[3]   # колонка D — «Сумма»
                assert not sum_cell.value
                assert sum_cell.fill.fgColor.rgb in ("00000000", None)
                date_cell = row[4]  # колонка E — «Дата» — расхождение есть, подсвечено
                assert date_cell.fill.fgColor.rgb == "00FFE0E0"
                return
        pytest.fail("Строка 'Целевые взносы ЮЛ' не найдена")

    def test_cel_ul_in_details(self, temp_data_dir: Path, output_file: Path):
        """Дневные расхождения Целевых взносов ЮЛ присутствуют в «Детализации»."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Детализация"]
        dates = self._details_map(ws).get("Целевые взносы ЮЛ", {})
        assert "10.02.2026" in dates
        assert "24.02.2026" in dates

    def test_details_group_header_top_and_subtotal_bottom(self, temp_data_dir: Path, output_file: Path):
        """Детализация: название группы сверху, итог расхождения по показателю — снизу."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Детализация"]
        all_rows = list(ws.iter_rows(min_row=2, values_only=True))

        # Заголовок группы «Страховой резерв» — только название, без сумм (B..E пусты)
        hdr_idx = next(
            (i for i, r in enumerate(all_rows) if r[0] == "Страховой резерв"
             and not any(r[1:])),
            None,
        )
        assert hdr_idx is not None, "Нет строки-заголовка группы «Страховой резерв»"

        # Следующая строка — дневная: без названия (A пусто), с датой и суммами
        nxt = all_rows[hdr_idx + 1]
        assert not nxt[0]                 # название показателя не повторяется
        assert nxt[1] == "17.02.2026"     # дата расхождения
        assert "15 000,00" in str(nxt[4])  # разница за день

        # Через строку — итог расхождения по показателю (снизу группы)
        subtotal = all_rows[hdr_idx + 2]
        assert subtotal[0] == "Итого расхождение по показателю"
        assert "15 000,00" in str(subtotal[4])

    def test_details_minus_not_parentheses(self, temp_data_dir: Path, output_file: Path):
        """Отрицательная разница в «Детализации» — со знаком «−», а не в скобках."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Детализация"]
        # Целевые взносы ЮЛ, 24.02.2026 — разница −7 000,00
        diffs = self._details_map(ws).get("Целевые взносы ЮЛ", {})
        assert diffs.get("24.02.2026") == "−7 000,00"
        assert "(" not in str(diffs.get("24.02.2026"))

    def test_pv_fl_period_totals(self, temp_data_dir: Path, output_file: Path):
        """Итоги за период по Пенсионным взносам ФЛ: БУ > ПУ ровно на 5 000."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Свод"]
        row = self._svod_row(ws, "Пенсионные взносы ФЛ")
        assert "1 106 706,60" in str(row[1])   # B — БУ
        assert "1 101 706,60" in str(row[2])    # C — ПУ
        assert "5 000,00" in str(row[3])        # D — расхождение по итогу

    def test_vyplata_both_sides_reconcile(self, temp_data_dir: Path, output_file: Path):
        """Выплаты заполняются с обеих сторон (ПУ — из Excel, БУ — из JSON НПО) и сверяются."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Свод"]
        row = self._svod_row(ws, "Выплата пенсии")
        assert row[1] and row[2]            # обе стороны заполнены
        assert "86 800,00" in str(row[1])   # B — БУ
        assert "86 800,00" in str(row[2])   # C — ПУ
        assert not row[3]                   # итоги равны → «Сумма» пустая (не 0,00)
        assert not row[4]                   # дат расхождений нет

    def test_negative_correction_no_discrepancy(self, temp_data_dir: Path, output_file: Path):
        """Целевые взносы ФЛ с отрицательной корректировкой: итоги равны, расхождения нет."""
        _run_pipeline(temp_data_dir, output_file)
        ws = openpyxl.load_workbook(output_file)["Свод"]
        row = self._svod_row(ws, "Целевые взносы ФЛ")
        assert row[1] == row[2]            # B=БУ == C=ПУ
        assert not row[3]                  # итоги равны → «Сумма» пустая (не 0,00)
        assert not row[4]                  # дат расхождений нет

    def test_llm_artifacts_created(self, temp_data_dir: Path, output_file: Path):
        """Pipeline создаёт обезличенное саммари (JSON) и промт рядом с отчётом."""
        import json
        _run_pipeline(temp_data_dir, output_file)
        summary = output_file.with_name("Саммари_БУ_ПУ.json")
        prompt = output_file.with_name("Промт_для_LLM.txt")
        assert summary.exists() and prompt.exists()
        data = json.loads(summary.read_text(encoding="utf-8"))
        assert data["обезличено"] is True
        assert data["итоги"]["показателей_с_расхождениями"] == 3
        # В саммари нет персональных данных
        assert "Иванов" not in summary.read_text(encoding="utf-8")

    def test_missing_dirs_creates_them_and_exits_0(self, tmp_path: Path):
        """Если директории данных отсутствуют, pipeline создаёт их и возвращает 0."""
        from npf_recon.config import Config
        from npf_recon.pipeline import run

        nonexistent = tmp_path / "nonexistent_data"
        config = Config(
            pu_dir=nonexistent / "ПУ",
            bu_dir=nonexistent / "БУ",
            output_file=tmp_path / "out.xlsx",
            period_label="Тест",
            eps=0.005,
        )
        exit_code = run(config)
        assert exit_code == 0
        assert (nonexistent / "ПУ").is_dir()
        assert (nonexistent / "БУ").is_dir()
